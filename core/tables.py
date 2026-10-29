"""Đọc file Excel (.xlsx) thành bảng chữ cho các tool phân tích. CSV thì trình duyệt tự đọc.

Bộ đọc nhanh: giải nén + quét XML trực tiếp (nhanh ~10× openpyxl). Gặp file lạ / lỗi thì tự
quay về openpyxl để luôn đọc đúng.
"""
from __future__ import annotations

import base64
import datetime
import io
import re
import zipfile
from html import unescape

from . import ToolError

MAX_ROWS = 300_000
OLE_MAGIC = bytes.fromhex("d0cf11e0")    # Excel có mật khẩu hoặc .xls cũ
NS_MAIN = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
NS_REL = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
# mã định dạng có sẵn của Excel là ngày / giờ
BUILTIN_DATE = set(range(14, 23)) | {27, 30, 36, 45, 46, 47, 50, 57}


def _num(f):
    """Cùng 1 quy tắc cho cả 2 bộ đọc: số nguyên ghi gọn, số quá lớn / lẻ ghi dạng chuẩn."""
    return str(int(f)) if f.is_integer() and abs(f) < 1e15 else repr(f)


def _cell(v):
    if v is None:
        return ""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, datetime.datetime):
        return v.date().isoformat() if v.time() == datetime.time(0) else v.isoformat(sep=" ", timespec="seconds")
    if isinstance(v, datetime.date):
        return v.isoformat()
    if isinstance(v, float):
        return _num(v)
    if isinstance(v, int):
        return _num(float(v)) if abs(v) >= 1e15 else str(v)
    return str(v)


def parse_xlsx(b64, sheet=None):
    return parse_xlsx_bytes(base64.b64decode(b64 or ""), sheet)


def parse_xlsx_bytes(data, sheet=None):
    if data[:4] == OLE_MAGIC:
        raise ToolError("File Excel đang có mật khẩu (hoặc là .xls cũ) — mở bằng Excel, bỏ mật khẩu / lưu thành .xlsx rồi thử lại.")
    try:
        return _fast_parse(data, sheet)
    except ToolError:
        raise
    except Exception:  # noqa: BLE001 — file lạ → để openpyxl xử lý
        return _openpyxl_parse(data, sheet)


# ------------------------------------------------------------------ bộ đọc nhanh
_COL_RE = re.compile(r"([A-Z]+)")
_TAG_T = re.compile(r"<(?:\w+:)?t(?:\s[^>]*)?>(.*?)</(?:\w+:)?t>", re.S)
_CELL_RE = re.compile(r"<(?:\w+:)?c\b(\s[^>]*?)?(?:/>|>(.*?)</(?:\w+:)?c>)", re.S)
_ROW_RE = re.compile(r"<(?:\w+:)?row\b[^>]*?(?:/>|>(.*?)</(?:\w+:)?row>)", re.S)
_ATTR_RE = re.compile(r'(\w+)="([^"]*)"')
_V_RE = re.compile(r"<(?:\w+:)?v>(.*?)</(?:\w+:)?v>", re.S)
_RPH_RE = re.compile(r"<(?:\w+:)?rPh\b.*?</(?:\w+:)?rPh>", re.S)


def _col_index(ref):
    """"AB12" → 27 (đếm từ 0). Duyệt ký tự thay vì regex cho nhanh."""
    n = 0
    for ch in ref:
        o = ord(ch)
        if 65 <= o <= 90:
            n = n * 26 + o - 64
        else:
            break
    return n - 1 if n else None


def _text_of(xml):
    """Ghép mọi <t> (kể cả rich text nhiều đoạn), bỏ phần phiên âm <rPh>."""
    if "<rPh" in xml or ":rPh" in xml:
        xml = _RPH_RE.sub("", xml)
    return unescape("".join(_TAG_T.findall(xml)))


def _serial_to_date(v, date1904):
    base = datetime.datetime(1904, 1, 1) if date1904 else datetime.datetime(1899, 12, 30)
    d = base + datetime.timedelta(days=float(v))
    return d.date().isoformat() if d.time() == datetime.time(0) else d.isoformat(sep=" ", timespec="seconds")


def _fast_parse(data, sheet):
    from xml.etree import ElementTree as ET

    zf = zipfile.ZipFile(io.BytesIO(data))
    names = set(zf.namelist())
    wb = ET.fromstring(zf.read("xl/workbook.xml"))
    date1904 = any((e.get("date1904") in ("1", "true")) for e in wb.iter(NS_MAIN + "workbookPr"))
    rels = ET.fromstring(zf.read("xl/_rels/workbook.xml.rels"))
    target = {r.get("Id"): r.get("Target") for r in rels}
    sheets = [(s.get("name"), target.get(s.get(NS_REL + "id"))) for s in wb.iter(NS_MAIN + "sheet")]
    if not sheets or any(t is None for _, t in sheets):
        raise ValueError("không đọc được danh sách sheet")
    sheet_names = [n for n, _ in sheets]

    shared = []
    if "xl/sharedStrings.xml" in names:
        ss = zf.read("xl/sharedStrings.xml").decode("utf-8")
        shared = [_text_of(si) for si in re.findall(r"<(?:\w+:)?si\b[^>]*>(.*?)</(?:\w+:)?si>", ss, re.S)]

    date_styles = set()
    if "xl/styles.xml" in names:
        st = ET.fromstring(zf.read("xl/styles.xml"))
        custom = {}
        for nf in st.iter(NS_MAIN + "numFmt"):
            code = re.sub(r'"[^"]*"|\[[^\]]*\]|\\.', "", nf.get("formatCode", "")).lower()
            custom[int(nf.get("numFmtId", "0"))] = bool(re.search(r"[dmyhs]", code)) and "general" not in code
        xfs = st.find(NS_MAIN + "cellXfs")
        for i, xf in enumerate(xfs if xfs is not None else []):
            fid = int(xf.get("numFmtId", "0"))
            if fid in BUILTIN_DATE or custom.get(fid):
                date_styles.add(str(i))

    def path_of(t):
        t = t.lstrip("/")
        return t if t.startswith("xl/") else "xl/" + t

    def read_sheet(path):
        xml = zf.read(path).decode("utf-8")
        matrix = []
        for row in _ROW_RE.finditer(xml):
            body = row.group(1) or ""
            cells = {}
            for c in _CELL_RE.finditer(body):
                a = c.group(1) or ""
                inner = c.group(2) or ""
                # đọc thuộc tính bằng find (nhanh hơn regex ~5×): r="B12" s="3" t="s"
                i = a.find('r="')
                col = _col_index(a[i + 3:a.find('"', i + 3)]) if i >= 0 else None
                if col is None:
                    col = max(cells) + 1 if cells else 0
                i = a.find(' t="')
                t = a[i + 4:a.find('"', i + 4)] if i >= 0 else "n"
                if t == "inlineStr":
                    cells[col] = _text_of(inner)
                    continue
                i = inner.find("<v>")
                if i >= 0:
                    raw = inner[i + 3:inner.find("</v>", i)]
                else:
                    vm = _V_RE.search(inner)
                    raw = vm.group(1) if vm else ""
                if "&" in raw:
                    raw = unescape(raw)
                if raw == "":
                    val = ""
                elif t == "s":
                    val = shared[int(raw)]
                elif t == "b":
                    val = "TRUE" if raw == "1" else "FALSE"
                elif t in ("str", "e"):
                    val = raw
                else:
                    i = a.find(' s="')
                    if i >= 0 and a[i + 4:a.find('"', i + 4)] in date_styles:
                        val = _serial_to_date(raw, date1904)
                    else:
                        val = _num(float(raw))
                cells[col] = val
            if not cells:
                continue
            out = [cells.get(i, "") for i in range(max(cells) + 1)]
            while out and out[-1] == "":
                out.pop()
            if any(x.strip() for x in out):
                matrix.append(out)
                if len(matrix) > MAX_ROWS:
                    raise ToolError("File quá lớn (> {} dòng) — lọc bớt rồi thử lại.".format(MAX_ROWS))
        return matrix

    paths = dict(sheets)
    if sheet in paths:
        return {"sheets": sheet_names, "sheet": sheet, "matrix": read_sheet(path_of(paths[sheet]))}
    for n in sheet_names:   # sheet đầu tiên có dữ liệu
        m = read_sheet(path_of(paths[n]))
        if len(m) > 1 or n == sheet_names[-1]:
            return {"sheets": sheet_names, "sheet": n, "matrix": m}
    raise ValueError("không có sheet")


# ------------------------------------------------------------------ dự phòng: openpyxl
def _openpyxl_parse(data, sheet=None):
    try:
        import openpyxl
    except ImportError as exc:
        raise ToolError("Thiếu thư viện openpyxl — chạy lại start.bat hoặc bấm 'Cài thư viện' ở trang Cài đặt.") from exc
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise ToolError("Không đọc được file Excel (chỉ hỗ trợ .xlsx): " + str(exc)) from exc
    names = wb.sheetnames
    chosen = sheet if sheet in names else None
    if chosen is None:
        for n in names:
            ws = wb[n]
            if ws.max_row and ws.max_row > 1:
                chosen = n
                break
        chosen = chosen or names[0]
    matrix = []
    for row in wb[chosen].iter_rows(values_only=True):
        cells = [_cell(v) for v in row]
        while cells and cells[-1] == "":
            cells.pop()
        if any(c.strip() for c in cells):
            matrix.append(cells)
        if len(matrix) > MAX_ROWS:
            raise ToolError("File quá lớn (> {} dòng) — lọc bớt rồi thử lại.".format(MAX_ROWS))
    wb.close()
    return {"sheets": names, "sheet": chosen, "matrix": matrix}
