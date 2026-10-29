"""Selftest lớp lõi + bảo mật: mã hoá bí mật, xuất Excel, và server thật chạy trên cổng ngẫu nhiên
(khoá phiên, header bảo mật, chống DNS rebinding, nén gzip, ETag, upload Excel, giữ kết nối)."""
from __future__ import annotations

import gzip
import http.client
import io
import json
import os
import sys
import tempfile
import zipfile
import threading

FAILS = []


def check(name, cond, extra=""):
    print("  [{}] {}{}".format("PASS" if cond else "FAIL", name, ("  -> " + str(extra)) if extra else ""))
    if not cond:
        FAILS.append(name)


def test_secrets():
    from core import secure
    print("\n1) Lưu bí mật đã mã hoá")
    d = tempfile.mkdtemp()
    p = os.path.join(d, "token.dat")
    secure.save_secret(p, '{"refresh_token": "rt-bí-mật-123"}')
    raw = open(p, "rb").read()
    if secure.dpapi_available():
        check("file token KHÔNG chứa chữ thường", b"refresh_token" not in raw and b"rt-" not in raw)
        check("đánh dấu đã mã hoá DPAPI", secure.is_encrypted(p))
    else:
        print("  [SKIP] máy không có DPAPI (không phải Windows)")
    check("giải mã lại đúng nội dung", secure.load_secret(p).decode("utf-8") == '{"refresh_token": "rt-bí-mật-123"}')
    plain = os.path.join(d, "token.json")
    open(plain, "w", encoding="utf-8").write('{"a": 1}')
    q = os.path.join(d, "token2.dat")
    check("file chữ thường bản cũ được mã hoá rồi xoá", secure.migrate_plain(plain, q) and not os.path.exists(plain)
          and secure.load_secret(q) == b'{"a": 1}')
    open(os.path.join(d, "bad.dat"), "wb").write(b"GDK1\x01rac")
    try:
        secure.load_secret(os.path.join(d, "bad.dat"))
        check("file bí mật hỏng → báo lỗi rõ", False)
    except secure.SecretError as exc:
        check("file bí mật hỏng → báo lỗi rõ", True, str(exc)[:60])
    check("xoá an toàn", secure.wipe(q) and not os.path.exists(q))


def test_exporter():
    import openpyxl
    from core import ToolError, exporter
    print("\n2) Xuất Excel")
    data = exporter.to_xlsx([{"name": "Change/Log", "headers": ["Tên", "Số"],
                              "rows": [['=HYPERLINK("http://x","bấm")', 5], ["@SUM(A1)", -3.5]]}])
    ws = openpyxl.load_workbook(io.BytesIO(data)).active
    check("tên sheet hợp lệ", ws.title == "Change_Log", ws.title)
    check("chuỗi bắt đầu bằng = / @ KHÔNG thành công thức", ws["A2"].data_type == "s" and ws["A3"].data_type == "s")
    check("số giữ kiểu số", ws["B3"].value == -3.5)
    enc = exporter.to_xlsx([{"name": "S", "headers": ["a"], "rows": [[1]]}], password="mat-khau-123")
    check("có mật khẩu → file Office mã hoá", enc[:4] == bytes.fromhex("d0cf11e0"))
    try:
        import msoffcrypto
        f = msoffcrypto.OfficeFile(io.BytesIO(enc)); f.load_key(password="mat-khau-123"); out = io.BytesIO(); f.decrypt(out)
        check("mở lại được bằng đúng mật khẩu", openpyxl.load_workbook(io.BytesIO(out.getvalue())).active["A2"].value == 1)
        g = msoffcrypto.OfficeFile(io.BytesIO(enc)); g.load_key(password="sai-mat-khau")
        try:
            g.decrypt(io.BytesIO()); check("sai mật khẩu bị từ chối", False)
        except Exception:  # noqa: BLE001
            check("sai mật khẩu bị từ chối", True)
    except ImportError:
        print("  [SKIP] chưa cài msoffcrypto-tool")
    try:
        exporter.to_xlsx([{"name": "S", "headers": ["a"], "rows": []}], password="123")
        check("mật khẩu < 8 ký tự bị từ chối", False)
    except ToolError:
        check("mật khẩu < 8 ký tự bị từ chối", True)


def test_fast_xlsx_reader():
    import datetime
    import time

    import openpyxl
    from core import exporter, tables
    print("\n3) Bộ đọc Excel nhanh = kết quả openpyxl")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Tổng hợp"
    ws.append(["Ngày", "Level", "Tên", "Bật", "Tỉ lệ", "Công thức"])
    ws.append([datetime.date(2026, 10, 1), 1, "Lê Văn Đạt & <Co>", True, 0.25, "=B2*2"])
    ws.append([datetime.datetime(2026, 10, 2, 13, 30), 2, None, False, 1.5, "=B3*2"])
    ws.append([None, None, None, None, None, None])
    ws["C5"] = "ô lẻ giữa hàng"
    ws.append([datetime.date(2026, 10, 3), 3, "  có dấu cách  ", None, -7, None])
    ws2 = wb.create_sheet("Sheet 2")
    ws2.append(["a", "b"])
    ws2.append(["x", 1e20])
    buf = io.BytesIO(); wb.save(buf); data = buf.getvalue()
    fast, slow = tables._fast_parse(data, None), tables._openpyxl_parse(data, None)
    check("sheet đầu: ngày, giờ, true/false, số, ô trống, ký tự đặc biệt", fast["matrix"] == slow["matrix"],
          [r for r in zip(fast["matrix"], slow["matrix"]) if r[0] != r[1]][:2])
    check("chọn sheet khác", tables._fast_parse(data, "Sheet 2")["matrix"] == tables._openpyxl_parse(data, "Sheet 2")["matrix"])
    wb.epoch = openpyxl.utils.datetime.CALENDAR_MAC_1904
    buf = io.BytesIO(); wb.save(buf)
    check("hệ ngày 1904 (Excel Mac cũ)", tables._fast_parse(buf.getvalue(), None)["matrix"][1][0] == "2026-10-01")
    rows = [[i, "Level %d" % i, "=HYPERLINK(\"x\")" if i % 7 == 0 else "ok", i * 1.5] for i in range(30000)]
    t0 = time.perf_counter(); big = exporter.to_xlsx([{"name": "Data", "headers": ["id", "name", "note", "val"], "rows": rows}]); t_w = time.perf_counter() - t0
    t0 = time.perf_counter(); f = tables._fast_parse(big, None); t_f = time.perf_counter() - t0
    t0 = time.perf_counter(); s2 = tables._openpyxl_parse(big, None); t_s = time.perf_counter() - t0
    check("file 30.000 dòng: kết quả giống hệt openpyxl", f["matrix"] == s2["matrix"])
    check("đọc nhanh hơn openpyxl ≥ 2×", t_s / max(t_f, 1e-6) >= 2, "{:.2f}s → {:.2f}s (×{:.1f}) · ghi {:.2f}s".format(t_s, t_f, t_s / max(t_f, 1e-6), t_w))
    check("file lỗi → tự quay về openpyxl / báo lỗi rõ", _raises(lambda: tables.parse_xlsx_bytes(b"PK\x03\x04hong")))


def _raises(fn):
    from core import ToolError
    try:
        fn()
        return False
    except ToolError:
        return True


def test_server():
    import server
    from core.users import Users
    print()
    print("4) Server: tài khoản, phiên, khoá phiên, header bảo mật, nhật ký sử dụng, nén, ETag")
    clock = [1_000_000.0]
    tmp = tempfile.mkdtemp()
    server.USERS = Users(tmp, clock=lambda: clock[0])       # thư mục tạm — không đụng tài khoản thật
    real_toolcfg = server.TOOLCFG_FILE
    httpd = server.make_server(0)
    port = httpd.server_address[1]
    threading.Thread(target=httpd.serve_forever, daemon=True).start()
    tok = server.SESSION_TOKEN
    cname = "gdk_sid_%d" % port
    jar = {}

    def req(method, path, body=None, headers=None, conn=None, auth=True):
        c = conn or http.client.HTTPConnection("127.0.0.1", port, timeout=10)
        h = {"Host": "127.0.0.1:%d" % port}
        if auth and jar.get("sid"):
            h["Cookie"] = "%s=%s" % (cname, jar["sid"])
        h.update(headers or {})
        c.request(method, path, body=body, headers=h)
        r = c.getresponse()
        data = r.read()
        if r.getheader("Content-Encoding") == "gzip":
            data = gzip.decompress(data)
        return r, data

    def auth_post(action, payload, headers=None, sid=None):
        h = {"Content-Type": "application/json", "X-GDK-Login": "1"}
        if sid:
            h["Cookie"] = "%s=%s" % (cname, sid)
        h.update(headers or {})
        return req("POST", "/api/auth/" + action, body=json.dumps(payload).encode(), headers=h, auth=False)

    def sid_of(r):
        c = r.getheader("Set-Cookie") or ""
        return c.split(";")[0].split("=", 1)[1] if c.startswith(cname + "=") else None

    def api(method, path, payload=None, sid=None):
        h = {"X-GDK-Token": tok, "Content-Type": "application/json"}
        if sid:
            h["Cookie"] = "%s=%s" % (cname, sid)
        r, data = req(method, "/api/" + path, body=json.dumps(payload).encode() if payload is not None else None,
                      headers=h, auth=sid is None)
        try:
            return r, json.loads(data or b"{}")
        except ValueError:
            return r, {}

    try:
        # ---------------- chưa có tài khoản: chỉ thấy trang đăng nhập / tạo quản trị viên
        r, html = req("GET", "/")
        check("chưa đăng nhập: '/' trả trang đăng nhập, KHÔNG lộ khoá phiên",
              b"login.js" in html and tok.encode() not in html)
        r1, _ = req("GET", "/api/tools", headers={"X-GDK-Token": tok})
        r2, _ = req("GET", "/assets/kit.js")
        r3, _ = req("GET", "/assets/login.js")
        check("chưa đăng nhập: API → 401 (kể cả có khoá phiên), file của tool bị chặn, file trang đăng nhập được tải",
              r1.status == 401 and r2.status == 401 and r3.status == 200)
        r, data = req("GET", "/api/auth/status")
        check("trạng thái: chưa có tài khoản", json.loads(data)["setup"] is False)
        r, _ = auth_post("setup", {"username": "admin", "password": "mat-khau-dai-123"}, headers={"X-GDK-Login": ""})
        check("tạo quản trị viên thiếu header chống CSRF → 403", r.status == 403)
        r, _ = auth_post("setup", {"username": "admin", "password": "ngan"})
        check("mật khẩu < 8 ký tự bị từ chối", r.status == 400)
        r, _ = auth_post("setup", {"username": "Tên Có Dấu", "password": "mat-khau-dai-123"})
        check("tên đăng nhập có dấu / dấu cách bị từ chối", r.status == 400)
        r, _ = auth_post("setup", {"username": "Admin", "name": "Dũng NT", "password": "mat-khau-dai-123"})
        jar["sid"] = sid_of(r)
        cookie = r.getheader("Set-Cookie") or ""
        check("tạo quản trị viên → đăng nhập luôn, cookie HttpOnly + SameSite=Strict, không ghi nhớ",
              r.status == 200 and jar["sid"] and "HttpOnly" in cookie and "SameSite=Strict" in cookie and "Max-Age" not in cookie)
        r, _ = auth_post("setup", {"username": "hacker", "password": "chiem-quyen-123"})
        check("đã có tài khoản → không ai tạo thêm quản trị viên qua trang đăng nhập", r.status == 400)
        raw = b"".join(open(os.path.join(tmp, f), "rb").read() for f in os.listdir(tmp) if f.startswith("gdk.db"))
        check("cơ sở dữ liệu chỉ chứa băm scrypt, không có mật khẩu thật", b"mat-khau-dai-123" not in raw and b"scrypt" in raw)
        admin_sid = jar["sid"]

        # ---------------- đã đăng nhập: các lớp cũ vẫn hoạt động
        r, html = req("GET", "/")
        check("đã đăng nhập: trang chủ có khoá phiên", tok.encode() in html and b"__GDK_TOKEN__" not in html)
        csp = r.getheader("Content-Security-Policy") or ""
        check("có CSP chặn script lạ + chặn iframe", "script-src 'self'" in csp and "frame-ancestors 'none'" in csp)
        check("X-Frame-Options / nosniff / no-referrer", r.getheader("X-Frame-Options") == "DENY"
              and r.getheader("X-Content-Type-Options") == "nosniff" and r.getheader("Referrer-Policy") == "no-referrer")
        check("trang chủ không cache (vì chứa khoá phiên)", r.getheader("Cache-Control") == "no-store")
        check("không có script inline trong trang", b"<script>" not in html)
        r, _ = req("GET", "/login")
        check("đã đăng nhập mà vào /login → chuyển về trang chủ", r.status == 302 and r.getheader("Location") == "/")
        r, _ = req("GET", "/api/tools")
        check("API thiếu khoá phiên → 401", r.status == 401)
        r, _ = req("GET", "/api/tools", headers={"X-GDK-Token": "doan-bua"})
        check("API khoá sai → 401", r.status == 401)
        r, me = api("GET", "tools")
        check("API đúng khoá → 200, biết người đang dùng", r.status == 200 and me["user"]["username"] == "admin"
              and me["user"]["role"] == "admin" and me["user"]["name"] == "Dũng NT")
        r, _ = req("GET", "/api/tools", headers={"X-GDK-Token": tok, "Cookie": cname + "=gia-mao"})
        check("cookie phiên giả → 401", r.status == 401)
        r, _ = req("GET", "/api/tools", headers={"X-GDK-Token": tok, "Host": "evil.example:%d" % port})
        check("Host lạ (DNS rebinding) → 403", r.status == 403)
        r, _ = req("POST", "/api/google/status", body=b"{}", headers={"X-GDK-Token": tok, "Content-Type": "text/plain"})
        check("POST sai Content-Type → 400", r.status == 400)
        conn = http.client.HTTPConnection("127.0.0.1", port, timeout=10)
        r, _ = req("POST", "/api/google/status", body=b'{"x":' + b"1" * 5000 + b"}", headers={"Content-Type": "application/json"}, conn=conn)
        conn2 = conn if not r.getheader("Connection") == "close" else http.client.HTTPConnection("127.0.0.1", port, timeout=10)
        r2, _ = req("GET", "/api/tools", headers={"X-GDK-Token": tok}, conn=conn2)
        check("request bị từ chối không làm hỏng request sau (đóng kết nối)", r.status == 401 and r2.status == 200)
        r, js = req("GET", "/assets/kit.js", headers={"Accept-Encoding": "gzip"})
        raw = open(os.path.join(server.WEB, "assets", "kit.js"), "rb").read()
        check("file tĩnh nén gzip và giải nén đúng", r.getheader("Content-Encoding") == "gzip" and js == raw)
        etag = r.getheader("ETag")
        r, _ = req("GET", "/assets/kit.js", headers={"If-None-Match": etag})
        check("ETag: tải lại không đổi → 304", r.status == 304)
        r, data = req("GET", "/..%2f..%2fserver.py")
        check("path traversal không lộ mã nguồn", b"SESSION_TOKEN" not in data)
        r, data = req("GET", "/api/tools/drive-clone/files/khong-co/changelog.csv?t=" + tok)
        check("tải file bằng ?t= (cho link tải) được chấp nhận", r.status == 400 and "Không tìm thấy" in data.decode())
        body = json.dumps({"name": "x", "sheets": [{"name": "S", "headers": ["a", "b"], "rows": [[1, "=1+1"]]}]}).encode()
        r, data = req("POST", "/api/export/xlsx", body=body, headers={"X-GDK-Token": tok, "Content-Type": "application/json"})
        check("xuất Excel qua API", r.status == 200 and data[:2] == b"PK" and "x.xlsx" in (r.getheader("Content-Disposition") or ""))
        r, data2 = req("POST", "/api/table/xlsx", body=data, headers={"X-GDK-Token": tok, "Content-Type": "application/octet-stream"})
        m = json.loads(data2) if r.status == 200 else {}
        check("đọc Excel gửi dạng byte", m.get("matrix") == [["a", "b"], ["1", "=1+1"]], m.get("matrix"))
        r, st = api("GET", "security/status")
        check("trạng thái bảo mật: quản trị viên thấy nhật ký + số tài khoản", r.status == 200 and st["admin"] and st["users"] == 1
              and "audit" in st)

        # ---------------- quản trị viên tạo tài khoản → người dùng phải đổi mật khẩu tạm
        print("   Tài khoản: tạo / mật khẩu tạm / quyền / khoá / xoá / nhật ký sử dụng")
        r, made = api("POST", "admin/users/create", {"username": "Lan.Anh", "name": "Lan Anh", "role": "user"})
        temp = made.get("password", "")
        check("tạo tài khoản → mật khẩu tạm 14 ký tự, tên đăng nhập chuẩn hoá chữ thường",
              r.status == 200 and len(temp) == 14 and made["user"]["username"] == "lan.anh" and made["user"]["must_change"])
        r, _ = api("POST", "admin/users/create", {"username": "lan.anh"})
        check("trùng tên đăng nhập → từ chối", r.status == 400)
        r, d = auth_post("login", {"username": "LAN.ANH", "password": temp})
        lan_sid = sid_of(r)
        check("đăng nhập bằng mật khẩu tạm (không phân biệt hoa thường tên) → bắt đổi mật khẩu", r.status == 200
              and json.loads(d)["must_change"] is True)
        r, _ = api("GET", "tools", sid=lan_sid)
        _, page = req("GET", "/", headers={"Cookie": cname + "=" + lan_sid}, auth=False)
        check("chưa đổi mật khẩu tạm: API bị chặn, trang chủ vẫn là trang đăng nhập", r.status == 401 and b"login.js" in page)
        r, _ = auth_post("change", {"old": temp, "new": "lan-anh-moi-789"}, sid=lan_sid)
        check("đổi mật khẩu tạm từ trang đăng nhập", r.status == 200)
        r, me = api("GET", "tools", sid=lan_sid)
        check("đổi xong → dùng được, vai trò user", r.status == 200 and me["user"]["role"] == "user")
        r, d = api("GET", "admin/users", sid=lan_sid)
        check("người dùng thường không gọi được API quản trị", r.status == 400 and "quản trị" in d.get("error", ""))
        r, _ = api("POST", "security/wipe", {"what": ["history"]}, sid=lan_sid)
        check("người dùng thường không xoá được lịch sử chung", r.status == 400)
        r, _ = api("POST", "usage", {"tool": "level-preview", "action": "open"}, sid=lan_sid)
        r2, _ = api("POST", "usage", {"tool": "level-preview", "action": "hack"}, sid=lan_sid)
        r3, _ = api("POST", "usage", {"tool": "khong-co", "action": "open"}, sid=lan_sid)
        check("ghi nhật ký mở tool; hành động / tool lạ bị từ chối", r.status == 200 and r2.status == 400 and r3.status == 400)
        r, d = api("GET", "admin/usage?uid=%d" % made["user"]["id"])
        acts = [x["action"] for x in d.get("rows", [])]
        check("quản trị viên xem nhật ký theo người: có đăng nhập + mở tool, kèm IP",
              r.status == 200 and "login" in acts and "open" in acts and d["rows"][0]["ip"] == "127.0.0.1", acts)
        r, d = api("GET", "admin/stats?days=7")
        check("thống kê theo tool / người", r.status == 200 and any(x["tool"] == "level-preview" for x in d["by_tool"])
              and any(x["username"] == "lan.anh" for x in d["by_user"]))
        logs = os.listdir(os.path.join(tmp, "logs"))
        text = open(os.path.join(tmp, "logs", logs[0]), encoding="utf-8-sig").read() if logs else ""
        check("có file nhật ký CSV theo tháng (mở bằng Excel)", len(logs) == 1 and logs[0].startswith("usage-")
              and "lan.anh" in text and "level-preview" in text, logs)
        check("tóm tắt thao tác không ghi mật khẩu / nội dung file",
              server._brief("games/save", {"password": "x", "content": "y", "name": "Ball"}) == "games/save · name=Ball")

        # ---------------- quản trị viên bật / tắt + sắp xếp tool
        print("   Tool: bật / tắt, thứ tự kéo thả, tự sắp theo mặc định")
        server.TOOLCFG_FILE = os.path.join(tmp, "tools.json")       # không đụng data/tools.json thật
        server._TOOLCFG["cfg"] = None
        default = list(server.TOOLS)
        hidden_default = [k for k, (m, _) in server.TOOLS.items() if m.get("hidden")]
        tid = "level-preview"
        r, d = api("GET", "tools", sid=lan_sid)
        check("mặc định: tool khai báo \"hidden\" (đang sửa) không gửi cho người dùng thường",
              r.status == 200 and not any(t["id"] in hidden_default for t in d["tools"]) and d["tools_auto"] is True)
        r, st = api("POST", "admin/tools/save", {"enabled": {tid: False}})
        r2, du = api("GET", "tools", sid=lan_sid)
        r3, da = api("GET", "tools")
        check("tắt tool → người dùng thường không thấy; quản trị viên vẫn thấy (kèm hidden)",
              r.status == 200 and all(t["id"] != tid for t in du["tools"])
              and any(t["id"] == tid and t["hidden"] for t in da["tools"]))
        r, d = api("GET", "tools/%s/games" % tid, sid=lan_sid)
        r2, _ = api("GET", "tools/%s/games" % tid)
        check("tool đang tắt: người dùng thường gọi chức năng → 403; quản trị viên vẫn thử được",
              r.status == 403 and "tạm tắt" in d.get("error", "") and r2.status == 200, (r.status, r2.status))
        r, _ = api("POST", "admin/tools/save", {"enabled": {tid: False}}, sid=lan_sid)
        check("người dùng thường không đổi được cấu hình tool", r.status == 400)
        order = default[::-1]
        r, st = api("POST", "admin/tools/save", {"order": order, "auto": False})
        r2, da = api("GET", "tools")
        check("kéo thả: thứ tự tuỳ chỉnh áp dụng cho danh sách tool",
              r.status == 200 and [t["id"] for t in da["tools"]] == order and st["auto"] is False)
        r, st = api("POST", "admin/tools/save", {"auto": True})
        r2, da = api("GET", "tools")
        check("bật 'tự sắp theo mặc định' → về thứ tự khai báo, thứ tự tuỳ chỉnh vẫn được giữ",
              [t["id"] for t in da["tools"]] == default and st["custom"] is True)
        api("POST", "admin/tools/save", {"auto": False})
        check("tắt lại → dùng lại thứ tự đã kéo", [t["id"] for t in api("GET", "tools")[1]["tools"]] == order)
        r, _ = api("POST", "admin/tools/save", {"order": order[:-1] + ["khong-co"]})
        r2, _ = api("POST", "admin/tools/save", {"order": order + [order[0]]})
        check("thứ tự có tool lạ / trùng → từ chối", r.status == 400 and r2.status == 400)
        server.TOOLS["zz-moi"] = ({"id": "zz-moi", "name": "Tool mới"}, None)
        try:
            ids = [t["id"] for t in api("GET", "tools")[1]["tools"]]
            check("tool mới thêm (chưa có trong thứ tự tuỳ chỉnh) tự chèn vào sau tool đứng trước nó",
                  ids.index("zz-moi") == ids.index(default[-1]) + 1, ids)
        finally:
            del server.TOOLS["zz-moi"]
        r, st = api("POST", "admin/tools/save", {"reset": True, "enabled": {tid: True}})
        check("về mặc định + bật lại", st["auto"] is True and st["custom"] is False
              and [t["id"] for t in st["tools"]] == default and all(t["enabled"] for t in st["tools"] if t["id"] == tid))
        check("cấu hình lưu ở thư mục dữ liệu (tools.json)", os.path.exists(server.TOOLCFG_FILE))

        # ---------------- Google: đã TẮT OAuth client loại Desktop app — Drive Clone chạy bằng Apps Script
        print("   Google: OAuth client (Desktop app) đã tắt · Drive Clone bản Apps Script")
        r, d = api("GET", "google/status", sid=lan_sid)
        r2, d2 = api("POST", "google/connect", {}, sid=lan_sid)
        r3, d3 = api("POST", "google/credentials", {"content": "{}"})
        r4, d4 = api("POST", "google/finish", {"url": "http://127.0.0.1:8765/?code=x"})
        check("kết nối Google qua OAuth client đã tắt: status báo tắt, connect / credentials / finish bị từ chối kể cả quản trị viên",
              r.status == 200 and d.get("disabled") is True and not d["connected"]
              and all(x.status == 400 and "Apps Script" in y.get("error", "") for x, y in ((r2, d2), (r3, d3), (r4, d4))))
        import tools.drive_clone as dclone
        real_dc_file = dclone._cfg_file
        dclone._cfg_file = lambda: os.path.join(tmp, "drive_clone.json")        # không đụng data/drive_clone.json thật
        try:
            r, d = api("GET", "tools/drive-clone/config", sid=lan_sid)
            check("Drive Clone: người dùng đọc được link web app + danh sách file Apps Script",
                  r.status == 200 and d["webapp_url"] == "" and d["admin"] is False and "Code.gs" in d["files"])
            r, d = api("POST", "tools/drive-clone/save", {"webapp_url": "https://script.google.com/macros/s/" + "A" * 40 + "/exec"}, sid=lan_sid)
            check("Drive Clone: người dùng thường KHÔNG đặt được link web app", r.status == 400 and "quản trị viên" in d.get("error", ""))
            r, d = api("POST", "tools/drive-clone/save", {"webapp_url": "https://evil.example/macros/s/xyz/exec"})
            check("Drive Clone: link không phải web app Apps Script bị từ chối", r.status == 400)
            good = "https://script.google.com/a/macros/iecgames.com/s/" + "AKfy" * 12 + "/exec"
            r, d = api("POST", "tools/drive-clone/save", {"webapp_url": good})
            r2, d2 = api("GET", "tools/drive-clone/config", sid=lan_sid)
            check("Drive Clone: quản trị viên lưu link (cả tên miền Workspace) → mọi người thấy", r.status == 200 and d2["webapp_url"] == good)
            r, data = req("GET", "/api/tools/drive-clone/files/appsscript.zip?t=" + tok)
            names = zipfile.ZipFile(io.BytesIO(data)).namelist() if r.status == 200 else []
            index = zipfile.ZipFile(io.BytesIO(data)).read("DriveClone_AppsScript/Index.html").decode("utf-8") if r.status == 200 else ""
            check("Drive Clone: tải bộ Apps Script .zip 2 file + hướng dẫn (không cần appsscript.json), Index.html gộp sẵn (không còn include)",
                  r.status == 200 and {n.split("/")[-1] for n in names} == {"Code.gs", "Index.html", "HUONG_DAN.txt"}
                  and "var DC = (function" in index and "include(" not in index and "<?" not in index, names)
            r, data = req("GET", "/api/tools/drive-clone/files/demo.html?t=" + tok)
            page = data.decode("utf-8", "replace")
            check("Drive Clone: bản Demo .html ghép đủ lõi + giao diện, không còn scriptlet Apps Script",
                  r.status == 200 and "var DC = (function" in page and "demoSeed" in page and "<?" not in page)
            r, data = req("GET", "/api/tools/drive-clone/files/appsscript/..%2f..%2f..%2fserver.py?t=" + tok)
            check("Drive Clone: không tải được file ngoài bộ Apps Script", r.status == 400 and b"SESSION_TOKEN" not in data)
        finally:
            dclone._cfg_file = real_dc_file

        # ---------------- chống dò mật khẩu: theo tên đăng nhập và theo máy
        codes = [auth_post("login", {"username": "lan.anh", "password": "sai-%d" % i})[0].status for i in range(5)]
        r, data = auth_post("login", {"username": "lan.anh", "password": "lan-anh-moi-789"})
        check("sai 5 lần → khoá tạm, kể cả nhập đúng cũng bị chặn", codes[:4] == [400] * 4 and codes[4] == 429 and r.status == 429,
              (codes, r.status))
        clock[0] += 31
        r, _ = auth_post("login", {"username": "lan.anh", "password": "lan-anh-moi-789"})
        check("hết thời gian khoá → đăng nhập đúng được (và xoá bộ đếm sai của máy)", r.status == 200)
        r, d = auth_post("login", {"username": "ai-do-khong-co", "password": "x" * 9})
        check("tên không tồn tại báo chung “sai tên hoặc mật khẩu” (không lộ tên nào có thật)",
              r.status == 400 and "Sai tên đăng nhập hoặc mật khẩu" in json.loads(d)["error"])
        for i in range(4):
            auth_post("login", {"username": "doan-ten-%d" % i, "password": "x" * 9})
        r, _ = auth_post("login", {"username": "admin", "password": "mat-khau-dai-123"})
        check("dò nhiều tên từ cùng 1 máy cũng bị khoá theo máy", r.status == 429)
        clock[0] += 61

        # ---------------- ghi nhớ / hết hạn
        r, _ = auth_post("login", {"username": "admin", "password": "mat-khau-dai-123", "remember": True})
        remembered = sid_of(r)
        check("'ghi nhớ' có Max-Age 7 ngày", r.status == 200 and "Max-Age=604800" in (r.getheader("Set-Cookie") or ""))
        restarted = Users(tmp, clock=lambda: clock[0])
        check("khởi động lại server: phiên vẫn còn (lưu trong cơ sở dữ liệu)", bool(restarted.check(remembered)))
        restarted.close()
        clock[0] += 4 * 3600 + 5
        check("phiên thường hết hạn sau 4 giờ không dùng; phiên ghi nhớ vẫn còn",
              not server.USERS.check(admin_sid) and server.USERS.check(remembered))
        admin_sid = jar["sid"] = remembered

        # ---------------- quản trị: khoá / đặt lại / xoá
        lid = made["user"]["id"]
        r, _ = auth_post("login", {"username": "lan.anh", "password": "lan-anh-moi-789"})
        lan_sid = sid_of(r)
        r, _ = api("POST", "admin/users/update", {"id": lid, "active": False})
        r2, d2 = auth_post("login", {"username": "lan.anh", "password": "lan-anh-moi-789"})
        check("khoá tài khoản → phiên đang mở bị đăng xuất, không đăng nhập được nữa",
              r.status == 200 and not server.USERS.check(lan_sid) and r2.status == 400 and "bị khoá" in json.loads(d2)["error"])
        r, _ = api("POST", "admin/users/update", {"id": 1, "active": False})
        r2, _ = api("POST", "admin/users/delete", {"id": 1})
        r3, _ = api("POST", "admin/users/update", {"id": 1, "role": "user"})
        check("quản trị viên không tự khoá / tự xoá / tự bỏ quyền của mình", r.status == 400 and r2.status == 400 and r3.status == 400)
        r, d = api("POST", "admin/users/update", {"id": lid, "active": True, "role": "admin"})
        r2, d2 = api("POST", "admin/users/update", {"id": lid, "role": "user"})
        check("mở khoá + đổi vai trò người khác", r.status == 200 and d["user"]["role"] == "admin" and d["user"]["active"]
              and r2.status == 200 and d2["user"]["role"] == "user")
        r, d = api("POST", "admin/users/reset", {"id": lid})
        r2, d2 = auth_post("login", {"username": "lan.anh", "password": d.get("password", "")})
        check("đặt lại mật khẩu → mật khẩu tạm mới, lại bắt đổi", r.status == 200 and r2.status == 200
              and json.loads(d2)["must_change"] is True)
        r, _ = api("POST", "admin/users/delete", {"id": lid})
        r2, _ = auth_post("login", {"username": "lan.anh", "password": d.get("password", "")})
        r3, d3 = api("GET", "admin/users")
        check("xoá tài khoản → không đăng nhập được, nhật ký cũ vẫn giữ",
              r.status == 200 and r2.status == 400 and len(d3["users"]) == 1
              and any(x["username"] == "lan.anh" for x in server.USERS.usage()["rows"]))

        # ---------------- đổi mật khẩu / đăng xuất / lệnh set-password
        r, _ = auth_post("login", {"username": "admin", "password": "mat-khau-dai-123"})
        other = sid_of(r)
        r, _ = api("POST", "auth/change", {"old": "sai-roi-123", "new": "mat-khau-moi-456"})
        check("đổi mật khẩu: sai mật khẩu hiện tại → từ chối", r.status == 400)
        r, _ = api("POST", "auth/change", {"old": "mat-khau-dai-123", "new": "mat-khau-moi-456"})
        check("đổi mật khẩu → phiên khác bị đăng xuất, phiên hiện tại giữ nguyên",
              r.status == 200 and not server.USERS.check(other) and server.USERS.check(remembered))
        r, _ = api("POST", "auth/logout", {})
        r2, _ = req("GET", "/api/tools", headers={"X-GDK-Token": tok})
        check("đăng xuất → cookie bị xoá, gọi API → 401",
              r.status == 200 and "Max-Age=0" in (r.getheader("Set-Cookie") or "") and r2.status == 401)
        server.USERS.set_password_cli("admin", "dat-lai-bang-lenh-789")
        check("lệnh set-password đặt lại được khi quên",
              auth_post("login", {"username": "admin", "password": "mat-khau-moi-456"})[0].status == 400
              and auth_post("login", {"username": "admin", "password": "dat-lai-bang-lenh-789"})[0].status == 200)
    finally:
        httpd.shutdown()
        server.USERS.close()
        server.TOOLCFG_FILE = real_toolcfg
        server._TOOLCFG["cfg"] = None


def test_accounts_extra():
    """Chuyển từ bản 1 mật khẩu, mạng nội bộ, Google theo từng người."""
    import base64
    import hashlib
    import server
    from core import google_auth
    from core.users import Users
    print()
    print("5) Chuyển bản cũ · mạng nội bộ · Google theo từng người")
    tmp = tempfile.mkdtemp()
    salt = b"0123456789abcdef"
    kdf = {"name": "scrypt", "n": 2 ** 14, "r": 8, "p": 1}
    h = hashlib.scrypt(b"mat-khau-cu-123", salt=salt, n=kdf["n"], r=8, p=1, dklen=32)
    with open(os.path.join(tmp, "auth.json"), "w", encoding="utf-8") as fh:
        json.dump({"v": 1, "kdf": kdf, "salt": base64.b64encode(salt).decode(), "hash": base64.b64encode(h).decode(),
                   "changed": "x"}, fh)
    u = Users(tmp)
    ok = u.login("admin", "mat-khau-cu-123", False, "127.0.0.1")
    check("bản 1 mật khẩu → tự thành tài khoản quản trị “admin”, đúng mật khẩu cũ",
          bool(ok["token"]) and u.by_username("admin")["role"] == "admin" and not os.path.exists(os.path.join(tmp, "auth.json"))
          and os.path.exists(os.path.join(tmp, "auth.json.migrated")))
    u.close()

    old = dict(server.NET)
    try:
        server.NET.update({"lan": False, "hosts": {"127.0.0.1", "localhost", "[::1]"}})
        check("chỉ máy này: máy khác bị chặn", server.client_allowed("127.0.0.1") and not server.client_allowed("192.168.1.20")
              and not server.host_allowed("192.168.1.5:7700"))
        server.NET.update({"lan": True, "hosts": {"127.0.0.1", "localhost", "[::1]", "may-chu", "192.168.1.5"}})
        check("mạng nội bộ: nhận máy 192.168.x / 10.x, chặn địa chỉ Internet",
              server.client_allowed("192.168.1.20") and server.client_allowed("10.0.0.7")
              and not server.client_allowed("8.8.8.8") and not server.client_allowed("::ffff:8.8.8.8"))
        check("mạng nội bộ: Host là tên / IP nội bộ được nhận, tên miền lạ bị chặn (DNS rebinding)",
              server.host_allowed("192.168.1.5:7700") and server.host_allowed("MAY-CHU:7700") and server.host_allowed("10.1.2.3")
              and not server.host_allowed("evil.example:7700") and not server.host_allowed("8.8.8.8:7700")
              and not server.host_allowed("192.168.1.5.evil.example"))
    finally:
        server.NET.clear(); server.NET.update(old)

    code, state = google_auth.parse_pasted("http://127.0.0.1:8765/?state=abc&code=4/0Ab-XYZ&scope=drive")
    check("Google: đọc mã từ địa chỉ dán vào", code == "4/0Ab-XYZ" and state == "abc")
    check("Google: dán thẳng mã cũng được; Google báo Huỷ → lỗi rõ ràng",
          google_auth.parse_pasted(" 4/0Ab ")[0] == "4/0Ab" and _raises(lambda: google_auth.parse_pasted("http://127.0.0.1:8765/?error=access_denied")))
    reg = google_auth.Registry()
    check("Google: mỗi tài khoản Tool Kit 1 token riêng", reg.get(1).token != reg.get(2).token and reg.get(1) is reg.get(1)
          and _raises(lambda: reg.get(None)))
    g = reg.get(7)
    g.pending = {"flow": None, "state": "s1", "exp": 0}
    check("Google: link đăng nhập quá hạn → yêu cầu kết nối lại", _raises(lambda: g.finish("http://x/?state=s1&code=c")))
    import time as _t
    g.pending = {"flow": None, "state": "s1", "exp": _t.time() + 60}
    check("Google: địa chỉ của lần đăng nhập khác (state lệch) bị từ chối", _raises(lambda: g.finish("http://x/?state=KHAC&code=c")))

    # quyền kết nối Google: quản trị viên luôn được; người dùng cần được bật — và luôn là Google của CHÍNH họ
    admin, lan, khoa = {"id": 1, "role": "admin"}, {"id": 5, "role": "user", "google": 1}, {"id": 6, "role": "user", "google": 0}
    check("Google: mỗi người dùng kết nối Google của chính mình khi được cấp quyền; chưa được cấp bị chặn",
          reg.for_user(admin) is reg.get(1) and reg.for_user(lan) is reg.get(5) and reg.for_user(lan) is not reg.get(1)
          and _raises(lambda: reg.for_user(khoa)) and _raises(lambda: reg.for_user(None)))


def test_single_instance():
    """Mở start.bat lần 2 khi bản cũ còn chạy: Windows từng cho 2 server cùng nghe 1 cổng → trình duyệt vẫn vào bản cũ."""
    import shutil
    import subprocess
    import time
    import urllib.request
    import server
    print()
    print("6) Chỉ 1 server trên 1 cổng: bản mới thay bản cũ")
    sample = chr(10).join([
        "  Proto  Local Address          Foreign Address        State           PID",
        "  TCP    127.0.0.1:7700         0.0.0.0:0              LISTENING       348308",
        "  TCP    127.0.0.1:7700         0.0.0.0:0              ĐANG NGHE       205264",
        "  TCP    127.0.0.1:7700         127.0.0.1:55301        ESTABLISHED     348308",
        "  TCP    127.0.0.1:17700        0.0.0.0:0              LISTENING       999",
        "  TCP    [::1]:7700             [::]:0                 LISTENING       42"])
    check("đọc netstat: chỉ lấy tiến trình đang nghe đúng cổng (kể cả Windows tiếng Việt / IPv6)",
          server.listening_pids(7700, sample) == [42, 205264, 348308], server.listening_pids(7700, sample))
    root = "D:/gd-toolkit"
    venv = '".venv/Scripts/python.exe" server.py'
    procs = {
        1: {"ProcessId": 1, "ParentProcessId": 0, "Name": "explorer.exe", "CommandLine": "explorer"},
        10: {"ProcessId": 10, "ParentProcessId": 1, "Name": "cmd.exe", "CommandLine": 'cmd.exe /c ""D:/gd-toolkit/start.bat" "'},
        11: {"ProcessId": 11, "ParentProcessId": 10, "Name": "python.exe", "CommandLine": venv},
        12: {"ProcessId": 12, "ParentProcessId": 11, "Name": "python.exe", "CommandLine": venv},
        20: {"ProcessId": 20, "ParentProcessId": 1, "Name": "cmd.exe", "CommandLine": 'cmd.exe /c ""E:/other/start.bat" "'},
        21: {"ProcessId": 21, "ParentProcessId": 20, "Name": "python.exe", "CommandLine": "python server.py"},
        30: {"ProcessId": 30, "ParentProcessId": 1, "Name": "python.exe", "CommandLine": "python server.py test"},
        31: {"ProcessId": 31, "ParentProcessId": 30, "Name": "python.exe", "CommandLine": "python server.py --port 9"},
    }
    check("tắt cả cửa sổ start.bat cũ của thư mục này", server.old_instance_tree(12, procs, root) == 10, server.old_instance_tree(12, procs, root))
    check("không đụng start.bat của thư mục khác", server.old_instance_tree(21, procs, root) == 21)
    check("không leo lên tiến trình cha khác lệnh (vd tool test)", server.old_instance_tree(31, procs, root) == 31)
    if os.name != "nt":
        check("(bỏ qua phần chạy thật — chỉ cần trên Windows)", True)
        return
    a = server.make_server(0)
    port = a.server_address[1]
    try:
        server.make_server(port)
        check("cổng độc quyền: server thứ 2 cùng cổng bị từ chối", False)
    except OSError:
        check("cổng độc quyền: server thứ 2 cùng cổng bị từ chối", True)
    finally:
        a.server_close()
    # chạy thật 2 server.py trên cùng 1 cổng (dữ liệu tạm) — bản 2 phải tắt bản 1 rồi nghe cổng thay
    tmp = tempfile.mkdtemp(prefix="gdk-inst-")
    here = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    env = dict(os.environ, PYTHONIOENCODING="utf-8", PYTHONUNBUFFERED="1")
    cmd = [sys.executable, os.path.join(here, "server.py"), "--no-open", "--port", str(port), "--data-dir", tmp]
    procs_ = []

    def alive():
        try:
            with urllib.request.urlopen("http://127.0.0.1:%d/api/auth/status" % port, timeout=2) as r:
                return r.status == 200
        except Exception:  # noqa: BLE001
            return False
    try:
        old = subprocess.Popen(cmd, cwd=here, env=env, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        procs_.append(old)
        for _ in range(60):
            if alive():
                break
            time.sleep(0.25)
        check("bản 1 chạy", alive())
        new = subprocess.Popen(cmd, cwd=here, env=env, stdout=subprocess.PIPE, stderr=subprocess.STDOUT)
        procs_.append(new)
        for _ in range(120):
            if old.poll() is not None and alive():
                break
            time.sleep(0.25)
        check("bản 2 tắt bản 1", old.poll() is not None)
        check("bản 2 đang phục vụ, chỉ còn 1 tiến trình nghe cổng", alive() and len(server.listening_pids(port)) == 1,
              server.listening_pids(port))
        check("bản 2 vẫn sống", new.poll() is None)
    finally:
        for p_ in procs_:
            if p_.poll() is None:
                subprocess.run(["taskkill", "/PID", str(p_.pid), "/T", "/F"], capture_output=True)
        out = procs_[-1].stdout.read().decode("utf-8", "replace") if len(procs_) > 1 and procs_[-1].stdout else ""
        if out:
            check("bản 2 báo đã tắt bản cũ", "Đã tắt GD Toolkit bản cũ" in out, out[-300:])
        shutil.rmtree(tmp, ignore_errors=True)


def main():
    from core import audit
    FAILS.clear()
    real_audit = audit.AUDIT
    audit.AUDIT = os.path.join(tempfile.mkdtemp(), "audit.log")   # test không được ghi vào nhật ký bảo mật thật
    try:
        return _run_all()
    finally:
        audit.AUDIT = real_audit


def _run_all():
    print("\n=== SELFTEST lõi & bảo mật ===")
    test_secrets()
    test_exporter()
    test_fast_xlsx_reader()
    test_server()
    test_accounts_extra()
    test_single_instance()
    print("\n=== KẾT QUẢ: {} ===\n".format("TẤT CẢ ĐẠT" if not FAILS else "HỎNG {} mục: {}".format(len(FAILS), FAILS)))
    return 0 if not FAILS else 1


if __name__ == "__main__":
    for _s in (sys.stdout, sys.stderr):
        try:
            _s.reconfigure(encoding="utf-8", errors="replace")
        except Exception:  # noqa: BLE001
            pass
    sys.exit(main())
