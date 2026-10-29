"""Xuất bảng ra Excel (.xlsx), có thể đặt MẬT KHẨU (mã hoá chuẩn Office / ECMA-376 Agile,
AES-256 — Excel hỏi mật khẩu khi mở). Dùng chung cho mọi tool."""
from __future__ import annotations

import io
import re

from . import ToolError

MAX_CELLS = 5_000_000
_RISKY = ("=", "+", "-", "@")
_BAD_XML = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f]")


def _safe_title(name, used):
    t = re.sub(r"[\[\]:*?/\\]", "_", str(name or "Sheet"))[:31] or "Sheet"
    base, k = t, 2
    while t in used:
        t = (base[:28] + " " + str(k))[:31]
        k += 1
    used.add(t)
    return t


def _cell_value(v):
    if v is None:
        return None
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, (int, float)):
        return v
    return str(v)


def to_xlsx(sheets, password=None):
    """sheets = [{"name", "headers": [...], "rows": [[...], ...]}] → bytes .xlsx"""
    try:
        import openpyxl
        from openpyxl.cell import WriteOnlyCell
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise ToolError("Thiếu thư viện openpyxl — chạy lại start.bat.") from exc
    if not sheets:
        raise ToolError("Không có dữ liệu để xuất.")
    total = sum(len(s.get("rows") or []) * max(1, len(s.get("headers") or [])) for s in sheets)
    if total > MAX_CELLS:
        raise ToolError("Bảng quá lớn để xuất Excel ({:,} ô) — hãy xuất CSV.".format(total))

    wb = openpyxl.Workbook(write_only=True)
    used = set()
    head_font, head_fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F3864")
    for sh in sheets:
        ws = wb.create_sheet(_safe_title(sh.get("name"), used))
        headers = [str(h) for h in sh.get("headers") or []]
        rows = sh.get("rows") or []
        widths = [len(h) for h in headers]
        for r in rows[:2000]:
            for i, v in enumerate(r[:len(headers)]):
                widths[i] = max(widths[i], min(60, len(str(v)) if v is not None else 0))
        for i, w in enumerate(widths):
            ws.column_dimensions[get_column_letter(i + 1)].width = max(8, min(60, w + 2))
        ws.freeze_panes = "A2"
        hdr = []
        for h in headers:
            c = WriteOnlyCell(ws, value=h)
            c.font, c.fill, c.alignment = head_font, head_fill, Alignment(vertical="center")
            hdr.append(c)
        ws.append(hdr)
        n = len(headers)
        for r in rows:
            out = []
            for v in list(r)[:n]:
                val = _cell_value(v)
                if isinstance(val, str):
                    if _BAD_XML.search(val):
                        val = _BAD_XML.sub("", val)          # ký tự điều khiển làm hỏng file Excel
                    if val[:1] in _RISKY:
                        c = WriteOnlyCell(ws, value=val)
                        c.data_type = "s"    # chuỗi bắt đầu bằng "=" KHÔNG được thành công thức (chống chèn công thức)
                        val = c
                out.append(val)          # ô thường: ghi thẳng giá trị (nhanh hơn tạo đối tượng ô)
            ws.append(out)
    buf = io.BytesIO()
    wb.save(buf)
    data = buf.getvalue()
    if password:
        data = encrypt_office(data, password)
    return data


def encrypt_office(data, password):
    if len(password) < 8:
        raise ToolError("Mật khẩu cần ít nhất 8 ký tự.")
    try:
        from msoffcrypto.format.ooxml import OOXMLFile
    except ImportError as exc:
        raise ToolError("Thiếu thư viện msoffcrypto-tool — chạy lại start.bat để cài.") from exc
    out = io.BytesIO()
    OOXMLFile(io.BytesIO(data)).encrypt(password, out)
    return out.getvalue()
